import requests
from bs4 import BeautifulSoup
import time
import random
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from datetime import datetime
import os
import re
import sys

BASE_URL = "https://audiobookbay.lu"
HEADERS = {
    "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
                  "(KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36",
    "Accept": "text/html,application/xhtml+xml,application/xml;q=0.9,*/*;q=0.8",
    "Accept-Language": "en-US,en;q=0.5",
    "Referer": "https://audiobookbay.lu/",
}

MAX_PAGES = int(os.environ.get("MAX_PAGES", "50"))
DELAY_MIN = float(os.environ.get("DELAY_MIN", "2.0"))
DELAY_MAX = float(os.environ.get("DELAY_MAX", "4.0"))
SCRAPE_DETAILS = os.environ.get("SCRAPE_DETAILS", "true").lower() == "true"


def get_page(url, retries=3):
    for attempt in range(retries):
        try:
            resp = requests.get(url, headers=HEADERS, timeout=30)
            resp.raise_for_status()
            print(f"  [OK] {url} (status {resp.status_code}, {len(resp.text)} chars)")
            return resp.text
        except Exception as e:
            print(f"  Attempt {attempt+1} failed for {url}: {e}")
            if attempt < retries - 1:
                time.sleep(random.uniform(3, 6))
    print(f"  [FAILED] {url}")
    return None


def debug_html_structure(html, page_num):
    """Print diagnostic info about what selectors exist on the page."""
    soup = BeautifulSoup(html, "html.parser")
    print(f"\n  --- DEBUG: Page {page_num} structure ---")
    print(f"  Title tag: {soup.title.string if soup.title else 'NONE'}")
    print(f"  div.post count: {len(soup.select('div.post'))}")
    print(f"  article count: {len(soup.select('article'))}")
    print(f"  .postWrapper count: {len(soup.select('.postWrapper'))}")
    print(f"  div.postContent count: {len(soup.select('div.postContent'))}")
    print(f"  div.post-content count: {len(soup.select('div.post-content'))}")
    print(f"  h2 tags total: {len(soup.select('h2'))}")
    print(f"  h2 a[href] count: {len(soup.select('h2 a[href]'))}")
    # Show all unique class names on divs to find the right container
    div_classes = set()
    for div in soup.select("div[class]"):
        for c in div.get("class", []):
            div_classes.add(c)
    # Filter to likely post containers
    likely = [c for c in div_classes if any(x in c.lower() for x in
              ["post", "book", "entry", "item", "card", "content", "article", "listing"])]
    print(f"  Likely container classes: {sorted(likely)}")
    # Show first h2 text to confirm we're on the right page
    h2s = soup.select("h2")
    for h2 in h2s[:3]:
        print(f"  Sample h2: {h2.get_text(strip=True)[:80]}")
    print(f"  --- END DEBUG ---\n")


def parse_detail_page(html, url):
    """Parse a single book's detail page for full metadata."""
    soup = BeautifulSoup(html, "html.parser")
    data = {}

    try:
        full_text = soup.get_text("\n")

        cat_match = re.search(r'Category:\s*(.+)', full_text)
        data["Categories"] = cat_match.group(1).strip() if cat_match else ""

        kw_match = re.search(r'(?:Keywords?|Tags?):\s*(.+)', full_text)
        data["Keywords"] = kw_match.group(1).strip() if kw_match else ""

        lang_match = re.search(r'Language:\s*(.+)', full_text)
        data["Language"] = lang_match.group(1).strip() if lang_match else ""

        fmt_match = re.search(r'Format:\s*([^\n/]+)', full_text)
        data["Format"] = fmt_match.group(1).strip() if fmt_match else ""

        br_match = re.search(r'Bitrate:\s*([^\n]+)', full_text)
        data["Bitrate"] = br_match.group(1).strip() if br_match else ""

        fs_match = re.search(r'File\s*Size:\s*([^\n]+)', full_text)
        data["File Size"] = fs_match.group(1).strip() if fs_match else ""

    except Exception as e:
        print(f"  [DETAIL PARSE ERROR] {e}")

    return data


def parse_listing_page(html, page_num=0):
    """Parse a listing page and return list of basic book dicts."""
    soup = BeautifulSoup(html, "html.parser")
    books = []

    # Run debug on first page so we can see what's there
    if page_num <= 1:
        debug_html_structure(html, page_num)

    # ── Strategy 1: standard div.post ──────────────────────────────
    posts = soup.select("div.post")

    # ── Strategy 2: article tags ────────────────────────────────────
    if not posts:
        posts = soup.select("article")

    # ── Strategy 3: common wrapper class names ──────────────────────
    if not posts:
        for selector in [".postWrapper", ".post-wrapper", ".entry", ".book-item",
                         ".listing-item", ".audiobook", "div.postContent",
                         "div.post-content", ".content-post"]:
            posts = soup.select(selector)
            if posts:
                print(f"  Found posts via selector: {selector}")
                break

    # ── Strategy 4: any div containing an /abss/ link ───────────────
    if not posts:
        print("  Trying strategy 4: divs containing /abss/ links")
        abss_links = soup.select("a[href*='/abss/']")
        print(f"  Found {len(abss_links)} /abss/ links on page")
        seen_urls = set()
        for a in abss_links:
            url = a.get("href", "")
            if not url.startswith("http"):
                url = BASE_URL + url
            if url in seen_urls:
                continue
            seen_urls.add(url)

            full_title = a.get_text(strip=True)
            if not full_title:
                continue

            # Walk up to find the nearest useful parent container
            parent = a.find_parent("div") or a.find_parent()
            post_text = parent.get_text("\n") if parent else ""

            books.append(_extract_book(full_title, url, post_text))

        return books

    # ── Parse found posts ────────────────────────────────────────────
    print(f"  Found {len(posts)} post containers")
    for post in posts:
        try:
            title_tag = post.select_one("h2 a") or post.select_one("h1 a") or post.select_one("a[href*='/abss/']")
            if not title_tag:
                continue

            full_title = title_tag.get_text(strip=True)
            url = title_tag.get("href", "")
            if url and not url.startswith("http"):
                url = BASE_URL + url

            post_text = post.get_text("\n")
            books.append(_extract_book(full_title, url, post_text))

        except Exception as e:
            print(f"  [PARSE ERROR] {e}")
            continue

    return books


def _extract_book(full_title, url, post_text):
    """Extract a book dict from a title, url, and surrounding text."""
    # Split title — author usually after last " - "
    if " - " in full_title:
        parts = full_title.rsplit(" - ", 1)
        title, author = parts[0].strip(), parts[1].strip()
    else:
        title, author = full_title, ""

    categories = keywords = language = format_ = bitrate = file_size = date_posted = ""

    for line in post_text.split("\n"):
        line = line.strip()
        if not line:
            continue
        if line.startswith("Category:"):
            categories = line.replace("Category:", "").strip()
        elif re.match(r'Keywords?:', line):
            keywords = re.sub(r'^Keywords?:\s*', '', line).strip()
        elif line.startswith("Language:"):
            language = line.replace("Language:", "").strip()
        elif line.startswith("Format:"):
            parts = line.split("/")
            format_ = parts[0].replace("Format:", "").strip()
            if len(parts) > 1 and "Bitrate:" in parts[1]:
                bitrate = parts[1].replace("Bitrate:", "").strip()
        elif line.startswith("Bitrate:") and not bitrate:
            bitrate = line.replace("Bitrate:", "").strip()
        elif re.match(r'File\s*Size:', line):
            file_size = re.sub(r'File\s*Size:\s*', '', line).strip()
        elif line.startswith("Posted:"):
            date_posted = line.replace("Posted:", "").strip()

    if not date_posted:
        dm = re.search(r'Posted:\s*(\d+\s+\w+\s+\d{4})', post_text)
        if dm:
            date_posted = dm.group(1).strip()

    return {
        "Title": title,
        "Author": author,
        "URL": url,
        "Categories": categories,
        "Keywords": keywords,
        "Language": language,
        "Format": format_,
        "Bitrate": bitrate,
        "File Size": file_size,
        "Date Posted": date_posted,
    }


def get_total_pages(html):
    soup = BeautifulSoup(html, "html.parser")
    max_page = 1
    for a in soup.select("a[href]"):
        href = a.get("href", "")
        if "/page/" in href:
            try:
                num = int(href.rstrip("/").split("/page/")[-1].rstrip("/"))
                if num > max_page:
                    max_page = num
            except:
                pass
    print(f"  Detected {max_page} total pages on site")
    return max_page


def save_xlsx(books, filename):
    wb = Workbook()
    ws = wb.active
    ws.title = "Audiobooks"

    header_font = Font(name="Arial", bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill("solid", start_color="1F4E79")
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left = Alignment(horizontal="left", vertical="center", wrap_text=True)
    thin = Side(style="thin", color="CCCCCC")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    columns = ["Title", "Author", "URL", "Categories", "Keywords", "Language",
               "Format", "Bitrate", "File Size", "Date Posted"]
    col_widths = [45, 25, 60, 30, 40, 12, 10, 10, 12, 15]

    for col_idx, (col_name, width) in enumerate(zip(columns, col_widths), 1):
        cell = ws.cell(row=1, column=col_idx, value=col_name)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center
        cell.border = border
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    ws.row_dimensions[1].height = 22
    ws.freeze_panes = "A2"

    fill_light = PatternFill("solid", start_color="EEF4FB")
    fill_white = PatternFill("solid", start_color="FFFFFF")

    for row_idx, book in enumerate(books, 2):
        fill = fill_light if row_idx % 2 == 0 else fill_white
        for col_idx, col_name in enumerate(columns, 1):
            val = book.get(col_name, "") or ""
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.font = Font(name="Arial", size=10)
            cell.fill = fill
            cell.border = border
            cell.alignment = left if col_name != "Date Posted" else center

    ws2 = wb.create_sheet("Summary")
    ws2["A1"] = "ABB Scrape Summary"
    ws2["A1"].font = Font(name="Arial", bold=True, size=14)
    ws2["A2"] = "Total Books Scraped:"
    ws2["B2"] = len(books)
    ws2["A3"] = "Scrape Date:"
    ws2["B3"] = datetime.now().strftime("%Y-%m-%d %H:%M")
    ws2["A4"] = "Source:"
    ws2["B4"] = BASE_URL
    ws2["A5"] = "Pages Scraped:"
    ws2["B5"] = MAX_PAGES if MAX_PAGES else "All"
    ws2.column_dimensions["A"].width = 22
    ws2.column_dimensions["B"].width = 25

    wb.save(filename)
    print(f"Saved {len(books)} books to {filename}")


def scrape():
    print(f"Fetching page 1: {BASE_URL}/")
    first_html = get_page(f"{BASE_URL}/")
    if not first_html:
        print("Could not reach site. Exiting.")
        sys.exit(1)  # Exit with error code so GitHub Actions marks the job as failed

    total_pages = get_total_pages(first_html)
    end_page = min(total_pages, MAX_PAGES) if MAX_PAGES else total_pages
    print(f"Will scrape pages 1 to {end_page}.")

    all_books = []

    for page in range(1, end_page + 1):
        if page == 1:
            html = first_html
        else:
            url = f"{BASE_URL}/page/{page}/"
            html = get_page(url)
            if not html:
                print(f"  Skipping page {page}")
                continue

        books = parse_listing_page(html, page_num=page)

        if SCRAPE_DETAILS and books:
            for book in books:
                if book["URL"] and not book.get("Categories"):
                    detail_html = get_page(book["URL"])
                    if detail_html:
                        detail_data = parse_detail_page(detail_html, book["URL"])
                        book.update({k: v for k, v in detail_data.items() if v})
                    time.sleep(random.uniform(0.5, 1.5))

        all_books.extend(books)
        print(f"Page {page}/{end_page}: {len(books)} books | Total: {len(all_books)}")

        if page < end_page:
            time.sleep(random.uniform(DELAY_MIN, DELAY_MAX))

    return all_books


if __name__ == "__main__":
    print("=" * 50)
    print("AudioBook Bay Scraper")
    print(f"MAX_PAGES={MAX_PAGES} | DETAIL_SCRAPE={SCRAPE_DETAILS}")
    print("=" * 50)

    books = scrape()

    if books:
        os.makedirs("output", exist_ok=True)
        timestamp = datetime.now().strftime("%Y%m%d_%H%M")
        filename = f"output/audiobookbay_{timestamp}.xlsx"
        save_xlsx(books, filename)
        print(f"Done. File: {filename}")
    else:
        print("No books scraped — check the DEBUG output above for HTML structure info.")
        # Save a dummy file so the artifact upload doesn't fail with a misleading error
        os.makedirs("output", exist_ok=True)
        filename = f"output/audiobookbay_{datetime.now().strftime('%Y%m%d_%H%M')}_EMPTY.xlsx"
        wb = Workbook()
        ws = wb.active
        ws["A1"] = "No books scraped — check Actions log for DEBUG output"
        wb.save(filename)
        print(f"Saved empty placeholder to {filename}")
        sys.exit(1)
